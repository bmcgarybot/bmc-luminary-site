#!/usr/bin/env python3
"""
Export Arizona Department of Education active school sites into BMC buyer census artifacts.

Source: ADE FY2027 Active LEAs and Schools spreadsheet.
Inventory research only; no outreach or provider activation.
"""
from __future__ import annotations

import argparse
from collections import Counter
import io
import json
from pathlib import Path
import re
import urllib.parse
import urllib.request

from openpyxl import load_workbook

DEFAULT_SOURCE = "https://www.azed.gov/sites/default/files/2026-10/FY27%20Active%20LEAs%20and%20Schools_0.xlsx"

MARKETS = [
    "Ash Fork",
    "Bullhead City",
    "Camp Verde",
    "Chino Valley",
    "Cordes Lakes",
    "Cottonwood",
    "Dewey-Humboldt",
    "Flagstaff",
    "Golden Valley",
    "Kingman",
    "Mayer",
    "Phoenix",
    "Prescott",
    "Prescott Valley",
    "Scottsdale",
    "Sedona",
    "Seligman",
    "Skull Valley",
    "Tucson",
    "Williamson Valley",
    "Winslow",
]

USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/156.0.0.0 Safari/537.36"


def clean(value):
    return " ".join(str(value or "").split()).strip()


def norm(value):
    return re.sub(r"[^a-z0-9]+", "", clean(value).lower())


def slug(value):
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def download(url):
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
            "Referer": "https://www.azed.gov/finance/local-education-agencies",
        },
    )
    with urllib.request.urlopen(req, timeout=90) as response:
        return response.read()


def find_header(ws):
    for row_idx in range(1, min(ws.max_row, 40) + 1):
        values = [clean(ws.cell(row_idx, c).value) for c in range(1, ws.max_column + 1)]
        normalized = [norm(v) for v in values]
        has_school = any("school" in v and ("name" in v or v == "school") for v in normalized)
        has_city = any(v == "city" or v.endswith("city") for v in normalized)
        if has_school and has_city:
            return row_idx, values
    raise RuntimeError(f"Could not find school/city header row in worksheet {ws.title!r}")


def choose_column(headers, candidates, contains=()):
    normalized = [norm(h) for h in headers]
    for candidate in candidates:
        key = norm(candidate)
        if key in normalized:
            return normalized.index(key) + 1
    for i, header in enumerate(normalized, 1):
        if all(term in header for term in contains):
            return i
    return None


def parse_workbook(blob):
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    rows = []
    diagnostics = []
    for ws in wb.worksheets:
        try:
            header_row, headers = find_header(ws)
        except RuntimeError:
            continue

        school_col = choose_column(headers, ["School Name", "School", "Site Name"], ("school",))
        city_col = choose_column(headers, ["City"], ("city",))
        address_col = choose_column(headers, ["Address", "School Address", "Site Address"], ("address",))
        zip_col = choose_column(headers, ["Zip", "Zip Code", "Postal Code"], ("zip",))
        lea_col = choose_column(headers, ["LEA Name", "District Name", "Charter Holder Name"], ("lea", "name"))
        if not lea_col:
            lea_col = choose_column(headers, ["District Name"], ("district",))
        entity_col = choose_column(headers, ["Entity ID", "School Entity ID", "Site Entity ID"], ("entity", "id"))
        status_col = choose_column(headers, ["Status", "School Status", "Entity Status"], ("status",))
        type_col = choose_column(headers, ["School Type", "Entity Type", "Type"], ("type",))
        phone_col = choose_column(headers, ["Phone", "Phone Number"], ("phone",))

        if not school_col or not city_col:
            continue

        sheet_count = 0
        for row_idx in range(header_row + 1, ws.max_row + 1):
            name = clean(ws.cell(row_idx, school_col).value)
            city = clean(ws.cell(row_idx, city_col).value)
            if not name or not city:
                continue
            status = clean(ws.cell(row_idx, status_col).value) if status_col else ""
            if status and status.lower() in {"inactive", "closed", "terminated"}:
                continue
            rows.append(
                {
                    "sheet": ws.title,
                    "row": row_idx,
                    "name": name,
                    "city": city,
                    "address": clean(ws.cell(row_idx, address_col).value) if address_col else "",
                    "zip": clean(ws.cell(row_idx, zip_col).value) if zip_col else "",
                    "lea": clean(ws.cell(row_idx, lea_col).value) if lea_col else "",
                    "entity_id": clean(ws.cell(row_idx, entity_col).value) if entity_col else "",
                    "status": status,
                    "school_type": clean(ws.cell(row_idx, type_col).value) if type_col else "",
                    "phone": clean(ws.cell(row_idx, phone_col).value) if phone_col else "",
                }
            )
            sheet_count += 1
        diagnostics.append({"sheet": ws.title, "header_row": header_row, "rows": sheet_count, "headers": headers})
    if not rows:
        raise RuntimeError("No active-school rows parsed from ADE workbook")
    return rows, diagnostics


def market_match(city, market):
    return clean(city).casefold() == market.casefold()


def build_records(rows, market, source_url, verified_date):
    records = []
    for item in rows:
        if not market_match(item["city"], market):
            continue
        anchor = item["entity_id"] or f"{item['sheet']}-row-{item['row']}"
        source = source_url + "#" + urllib.parse.quote(str(anchor), safe="-_.")
        note_parts = ["Arizona Department of Education FY2027 active LEA/school site"]
        if item["lea"]:
            note_parts.append(f"LEA: {item['lea']}")
        if item["school_type"]:
            note_parts.append(f"type: {item['school_type']}")
        if item["address"]:
            note_parts.append(f"address: {item['address']}, {market}, AZ {item['zip']}".strip())
        if item["entity_id"]:
            note_parts.append(f"entity ID {item['entity_id']}")
        records.append(
            {
                "side": "buyer",
                "market": market,
                "category_key": "schools_education",
                "category_label": "schools education childcare",
                "business_name": item["name"],
                "url": source,
                "domain": "azed.gov",
                "phone": item["phone"] or None,
                "email": None,
                "official": False,
                "official_registry": True,
                "verified_date": verified_date,
                "address": item["address"] or None,
                "zip": item["zip"] or None,
                "lea": item["lea"] or None,
                "entity_id": item["entity_id"] or None,
                "school_type": item["school_type"] or None,
                "note": " | ".join(note_parts),
            }
        )
    records.sort(key=lambda r: (r["business_name"].casefold(), r.get("entity_id") or ""))
    return records


def write_artifact(path, market, records, source_url, verified_date):
    payload = {
        "generated_at": verified_date + "T23:59:00Z",
        "purpose": f"Assistant-led {market} buyer census from Arizona Department of Education FY2027 active school-site records.",
        "safety": "Inventory seed only. Does not promote buyers/vendors and does not stage or send outreach.",
        "source": source_url,
        "count": len(records),
        "records": records,
    }
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--source-url", default=DEFAULT_SOURCE)
    args = parser.parse_args()

    blob = download(args.source_url)
    rows, diagnostics = parse_workbook(blob)
    args.output_dir.mkdir(parents=True, exist_ok=True)

    counts = {}
    files = []
    for market in MARKETS:
        records = build_records(rows, market, args.source_url, args.date)
        if not records:
            continue
        path = args.output_dir / f"assistant-research-bootstrap-{args.date}-ade-schools-{slug(market)}-08.json"
        write_artifact(path, market, records, args.source_url, args.date)
        counts[market] = len(records)
        files.append(str(path))

    if sum(counts.values()) < 100:
        raise RuntimeError(f"ADE market school export unexpectedly small: {sum(counts.values())} records")

    print(json.dumps({
        "source_rows": len(rows),
        "market_records": sum(counts.values()),
        "by_market": counts,
        "files": files,
        "worksheets": diagnostics,
    }, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
